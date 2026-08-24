"""Orchestration service for Build vs Buy evaluations.

Coordinates between scoring, TCO calculation, SharePoint persistence,
and audit logging.

Requirements: 1.1, 1.2, 1.3, 1.4, 1.5, 2.1, 2.3, 2.5, 2.6,
             4.1, 4.2, 4.5, 5.1, 6.6, 10.1, 10.2, 10.3, 10.4, 10.5, 12.1, 12.2
"""

import io
import logging
import uuid
from datetime import datetime, timezone
from typing import Any, Optional

from openpyxl import Workbook, load_workbook

from app.models.schemas import (
    BvBEvaluationResponse,
    DimensionScoreResponse,
    EvaluationCreateRequest,
    EvaluationStatus,
    OptionCreateRequest,
    OptionResponse,
    OptionType,
    ScoreResponse,
    TCOResultResponse,
    TCOSubmitRequest,
)
from app.services.sharepoint_client import (
    FOLDER_BUILD_VS_BUY,
    FOLDER_BUILD_VS_BUY_AUDIT,
    FOLDER_BUILD_VS_BUY_EVALUATIONS,
    SharePointClient,
    SharePointError,
)
from app.services import scoring_matrix, tco_calculator

logger = logging.getLogger(__name__)

# File paths
EVAL_INDEX_FILE = f"{FOLDER_BUILD_VS_BUY}/EvalIndex.xlsx"
EVAL_INDEX_SHEET = "Sheet1"
EVAL_INDEX_HEADERS = [
    "EvaluationId", "OpportunityName", "BusinessUnit",
    "Status", "OwnerName", "CreatedAt", "UpdatedAt", "FileName",
]

# Audit file config
AUDIT_SHEET = "Sheet1"
AUDIT_HEADERS = [
    "AuditId", "Timestamp", "EventType", "EvaluationId", "UserEmail", "Details",
]

# Per-evaluation workbook sheet names
SHEET_METADATA = "Metadata"
SHEET_OPTIONS = "Options"
SHEET_SCORES = "Scores"
SHEET_TCO = "TCO"


def _evaluation_file(evaluation_id: str) -> str:
    """Return the SharePoint path for a per-evaluation workbook (UUID fallback)."""
    return f"{FOLDER_BUILD_VS_BUY_EVALUATIONS}/{evaluation_id}.xlsx"


def _make_readable_filename(opportunity_name: str) -> str:
    """Create a readable, unique filename from the opportunity name + timestamp.

    Example: 'CMS Platform Selection' → 'CMS-Platform-Selection_20260628_143021.xlsx'
    Timestamp ensures uniqueness even if names repeat.
    """
    import re
    clean = re.sub(r'[^a-zA-Z0-9\s]', '', opportunity_name).strip()
    clean = re.sub(r'\s+', '-', clean)[:40]
    if not clean:
        clean = "Evaluation"
    ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    return f"{clean}_{ts}.xlsx"


async def _resolve_eval_file(evaluation_id: str, sp: SharePointClient) -> str:
    """Resolve the actual file path for an evaluation from the index."""
    try:
        rows = await sp.read_workbook(EVAL_INDEX_FILE, sheet=EVAL_INDEX_SHEET)
        for row in rows[1:]:
            if row and str(row[0]) == evaluation_id and len(row) > 7 and row[7]:
                return f"{FOLDER_BUILD_VS_BUY_EVALUATIONS}/{row[7]}"
    except Exception:
        pass
    raise SharePointError(f"Evaluation file not found for ID: {evaluation_id}")


def _now_iso() -> str:
    """Return current UTC time in ISO format."""
    return datetime.now(timezone.utc).isoformat()



# ---------------------------------------------------------------------------
# Helper: Create evaluation Excel workbook in memory
# ---------------------------------------------------------------------------


def _create_evaluation_workbook(metadata: dict) -> bytes:
    """Create a new per-evaluation Excel workbook with all sheets and metadata.

    Args:
        metadata: Dict of metadata key-value pairs to write to the Metadata sheet.

    Returns:
        Excel file bytes.
    """
    wb = Workbook()

    # Metadata sheet
    ws_meta = wb.active
    ws_meta.title = SHEET_METADATA
    ws_meta.append(["Key", "Value"])
    for key, value in metadata.items():
        ws_meta.append([key, str(value) if value is not None else ""])

    # Options sheet
    ws_options = wb.create_sheet(SHEET_OPTIONS)
    ws_options.append(["OptionId", "Name", "Type", "VendorName", "Description"])

    # Scores sheet
    ws_scores = wb.create_sheet(SHEET_SCORES)
    ws_scores.append(["OptionId", "Dimension", "Score", "Source"])

    # TCO sheet
    ws_tco = wb.create_sheet(SHEET_TCO)
    ws_tco.append(["OptionId", "Category", "Year", "Amount"])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Helper: Create or ensure EvalIndex.xlsx exists
# ---------------------------------------------------------------------------


async def _ensure_eval_index(sp: SharePointClient) -> None:
    """Ensure the EvalIndex.xlsx file exists, creating it if needed."""
    if not await sp.file_exists(EVAL_INDEX_FILE):
        await sp._create_excel_file(EVAL_INDEX_FILE, EVAL_INDEX_HEADERS)
        logger.info(f"Created evaluation index file: {EVAL_INDEX_FILE}")


# ---------------------------------------------------------------------------
# Helper: Audit logging
# ---------------------------------------------------------------------------


def _audit_file_path() -> str:
    """Return the SharePoint path for the current month's BvB audit file."""
    now = datetime.now(timezone.utc)
    return f"{FOLDER_BUILD_VS_BUY_AUDIT}/BvBAudit_{now.strftime('%Y-%m')}.xlsx"


async def _ensure_audit_file(sp: SharePointClient) -> str:
    """Ensure the current month's BvB audit file exists, creating if needed.

    Returns:
        The path to the audit file.
    """
    audit_file = _audit_file_path()
    if not await sp.file_exists(audit_file):
        await sp._create_excel_file(audit_file, AUDIT_HEADERS)
        logger.info(f"Created BvB audit file: {audit_file}")
    return audit_file


async def _log_audit_event(
    sp: SharePointClient,
    event_type: str,
    evaluation_id: str,
    user_email: str,
    details: str,
) -> None:
    """Log an audit event to the monthly BvB audit file.

    Does not raise exceptions — failures are logged but do not break the main flow.

    Args:
        sp: SharePoint client instance.
        event_type: Type of event (e.g. "evaluation_created").
        evaluation_id: The evaluation UUID.
        user_email: Email of the acting user.
        details: Free-text description of the event.
    """
    try:
        audit_file = await _ensure_audit_file(sp)
        audit_id = str(uuid.uuid4())
        timestamp = _now_iso()
        row = [audit_id, timestamp, event_type, evaluation_id, user_email, details]
        await sp.append_row(audit_file, AUDIT_SHEET, row)
        logger.debug(f"Audit logged: {event_type} for evaluation {evaluation_id}")
    except Exception as exc:
        logger.warning(f"Failed to log BvB audit event: {exc}")



# ---------------------------------------------------------------------------
# 1. create_evaluation
# ---------------------------------------------------------------------------


async def create_evaluation(
    request: EvaluationCreateRequest,
    user_display_name: str,
    user_email: str,
    sp: SharePointClient,
) -> BvBEvaluationResponse:
    """Create a new Build vs Buy evaluation.

    Validates intake fields, generates a UUID, persists to SharePoint, and
    updates the evaluation index.

    Args:
        request: The evaluation creation request.
        user_display_name: Display name of the creating user.
        user_email: Email of the creating user.
        sp: SharePoint client instance.

    Returns:
        BvBEvaluationResponse with the created evaluation details.

    Raises:
        ValueError: If validation fails (empty name, budget_max < budget_min).
    """
    # Validate
    if not request.opportunity_name or not request.opportunity_name.strip():
        raise ValueError("opportunity_name must be non-empty")
    if request.budget_max < request.budget_min:
        raise ValueError("budget_max must be >= budget_min")

    evaluation_id = str(uuid.uuid4())
    now = _now_iso()

    metadata = {
        "evaluation_id": evaluation_id,
        "opportunity_name": request.opportunity_name,
        "problem_statement": request.problem_statement,
        "business_unit": request.business_unit,
        "budget_min": request.budget_min,
        "budget_max": request.budget_max,
        "target_decision_date": request.target_decision_date,
        "owner_display_name": user_display_name,
        "owner_email": user_email,
        "status": EvaluationStatus.DRAFT.value,
        "created_at": now,
        "updated_at": now,
    }

    readable_name = _make_readable_filename(request.opportunity_name)
    eval_file = f"{FOLDER_BUILD_VS_BUY_EVALUATIONS}/{readable_name}"
    metadata["file_name"] = readable_name
    workbook_bytes = _create_evaluation_workbook(metadata)
    await sp._upload_file(eval_file, workbook_bytes)

    # Update evaluation index
    await _ensure_eval_index(sp)
    index_row = [
        evaluation_id,
        request.opportunity_name,
        request.business_unit,
        EvaluationStatus.DRAFT.value,
        user_display_name,
        now,
        now,
        readable_name,
    ]
    await sp.append_row(EVAL_INDEX_FILE, EVAL_INDEX_SHEET, index_row)

    logger.info(f"Created evaluation {evaluation_id}: '{request.opportunity_name}'")

    # Log audit event (non-blocking, won't break main flow)
    try:
        await _log_audit_event(
            sp=sp,
            event_type="evaluation_created",
            evaluation_id=evaluation_id,
            user_email=user_email,
            details=f"Created evaluation '{request.opportunity_name}' for {request.business_unit}",
        )
    except Exception as exc:
        logger.warning(f"Audit logging failed for evaluation {evaluation_id}: {exc}")

    return BvBEvaluationResponse(
        id=evaluation_id,
        opportunity_name=request.opportunity_name,
        problem_statement=request.problem_statement,
        business_unit=request.business_unit,
        budget_min=request.budget_min,
        budget_max=request.budget_max,
        target_decision_date=request.target_decision_date,
        owner_display_name=user_display_name,
        status=EvaluationStatus.DRAFT,
        created_at=now,
        updated_at=now,
    )



# ---------------------------------------------------------------------------
# 2. load_evaluation
# ---------------------------------------------------------------------------


async def load_evaluation(evaluation_id: str, sp: SharePointClient) -> dict:
    """Load a complete evaluation from its SharePoint workbook.

    Reads and parses all sheets (Metadata, Options, Scores, TCO).

    Args:
        evaluation_id: UUID of the evaluation to load.
        sp: SharePoint client instance.

    Returns:
        Dict with keys: metadata, options, scores, tco.

    Raises:
        SharePointError: If the evaluation file cannot be found/read.
    """
    eval_file = await _resolve_eval_file(evaluation_id, sp)
    file_bytes = await sp._download_file(eval_file)
    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)

    result: dict[str, Any] = {
        "metadata": {},
        "options": [],
        "scores": [],
        "tco": [],
    }

    # Parse Metadata sheet (Key-Value pairs)
    if SHEET_METADATA in wb.sheetnames:
        ws = wb[SHEET_METADATA]
        rows = list(ws.iter_rows(values_only=True))
        for row in rows[1:]:  # skip header
            if row and len(row) >= 2 and row[0]:
                result["metadata"][str(row[0])] = row[1] if row[1] is not None else ""

    # Parse Options sheet
    if SHEET_OPTIONS in wb.sheetnames:
        ws = wb[SHEET_OPTIONS]
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) > 1:
            headers = [str(h) if h else "" for h in rows[0]]
            for row in rows[1:]:
                if row and row[0]:
                    option = {}
                    for i, header in enumerate(headers):
                        option[header] = row[i] if i < len(row) and row[i] is not None else ""
                    result["options"].append(option)

    # Parse Scores sheet
    if SHEET_SCORES in wb.sheetnames:
        ws = wb[SHEET_SCORES]
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) > 1:
            headers = [str(h) if h else "" for h in rows[0]]
            for row in rows[1:]:
                if row and row[0]:
                    score = {}
                    for i, header in enumerate(headers):
                        score[header] = row[i] if i < len(row) and row[i] is not None else ""
                    result["scores"].append(score)

    # Parse TCO sheet
    if SHEET_TCO in wb.sheetnames:
        ws = wb[SHEET_TCO]
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) > 1:
            headers = [str(h) if h else "" for h in rows[0]]
            for row in rows[1:]:
                if row and row[0]:
                    tco_row = {}
                    for i, header in enumerate(headers):
                        tco_row[header] = row[i] if i < len(row) and row[i] is not None else ""
                    result["tco"].append(tco_row)

    wb.close()
    return result


# ---------------------------------------------------------------------------
# 3. save_evaluation_metadata
# ---------------------------------------------------------------------------


async def save_evaluation_metadata(
    evaluation_id: str, metadata: dict, sp: SharePointClient
) -> None:
    """Write metadata to the Metadata sheet of an evaluation workbook.

    Overwrites the existing metadata in the workbook.

    Args:
        evaluation_id: UUID of the evaluation.
        metadata: Dict of key-value pairs to write.
        sp: SharePoint client instance.
    """
    eval_file = await _resolve_eval_file(evaluation_id, sp)
    file_bytes = await sp._download_file(eval_file)
    wb = load_workbook(io.BytesIO(file_bytes))

    # Clear and rewrite Metadata sheet
    if SHEET_METADATA in wb.sheetnames:
        ws = wb[SHEET_METADATA]
        wb.remove(ws)
    ws = wb.create_sheet(SHEET_METADATA, 0)
    ws.append(["Key", "Value"])
    for key, value in metadata.items():
        ws.append([key, str(value) if value is not None else ""])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    await sp._upload_file(eval_file, buffer.getvalue())


async def update_index_status(
    evaluation_id: str, new_status: str, sp: SharePointClient
) -> None:
    """Update the status column for an evaluation in EvalIndex.xlsx."""
    try:
        if not await sp.file_exists(EVAL_INDEX_FILE):
            return

        file_bytes = await sp._download_file(EVAL_INDEX_FILE)
        wb = load_workbook(io.BytesIO(file_bytes))
        ws = wb.active

        for row in ws.iter_rows(min_row=2, values_only=False):
            if row[0].value == evaluation_id:
                row[3].value = new_status  # Status column
                row[6].value = _now_iso()  # UpdatedAt column
                break

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        await sp._upload_file(EVAL_INDEX_FILE, buffer.getvalue())
        wb.close()
    except Exception as exc:
        logger.warning(f"Failed to update index status for {evaluation_id}: {exc}")



# ---------------------------------------------------------------------------
# 4. list_evaluations
# ---------------------------------------------------------------------------


async def list_evaluations(sp: SharePointClient) -> list[BvBEvaluationResponse]:
    """List all evaluations from the EvalIndex.xlsx file.

    Args:
        sp: SharePoint client instance.

    Returns:
        List of BvBEvaluationResponse objects.
    """
    if not await sp.file_exists(EVAL_INDEX_FILE):
        return []

    rows = await sp.read_workbook(EVAL_INDEX_FILE, sheet=EVAL_INDEX_SHEET)
    if not rows or len(rows) < 2:
        return []

    evaluations: list[BvBEvaluationResponse] = []
    # First row is headers
    for row in rows[1:]:
        if not row or not row[0]:
            continue
        try:
            evaluations.append(
                BvBEvaluationResponse(
                    id=str(row[0]),
                    opportunity_name=str(row[1]) if len(row) > 1 and row[1] else "",
                    problem_statement="",  # Not stored in index
                    business_unit=str(row[2]) if len(row) > 2 and row[2] else "",
                    budget_min=0.0,  # Not stored in index
                    budget_max=0.0,  # Not stored in index
                    target_decision_date="",  # Not stored in index
                    owner_display_name=str(row[4]) if len(row) > 4 and row[4] else "",
                    status=EvaluationStatus(str(row[3])) if len(row) > 3 and row[3] else EvaluationStatus.DRAFT,
                    created_at=str(row[5]) if len(row) > 5 and row[5] else "",
                    updated_at=str(row[6]) if len(row) > 6 and row[6] else "",
                )
            )
        except (ValueError, IndexError) as exc:
            logger.warning(f"Skipping invalid index row: {exc}")
            continue

    return evaluations


# ---------------------------------------------------------------------------
# 5. add_option
# ---------------------------------------------------------------------------


async def add_option(
    evaluation_id: str,
    option: OptionCreateRequest,
    sp: SharePointClient,
) -> OptionResponse:
    """Add an option to an evaluation.

    Validates that the evaluation has fewer than 5 options already.

    Args:
        evaluation_id: UUID of the evaluation.
        option: The option creation request.
        sp: SharePoint client instance.

    Returns:
        OptionResponse with the new option details.

    Raises:
        ValueError: If the evaluation already has 5 options.
    """
    data = await load_evaluation(evaluation_id, sp)
    if len(data["options"]) >= 5:
        raise ValueError("Maximum of 5 options allowed per evaluation")

    option_id = str(uuid.uuid4())

    # Append to Options sheet
    eval_file = await _resolve_eval_file(evaluation_id, sp)
    await sp.append_row(
        eval_file,
        SHEET_OPTIONS,
        [
            option_id,
            option.name,
            option.type.value,
            option.vendor_name or "",
            option.description or "",
        ],
    )

    logger.info(f"Added option {option_id} to evaluation {evaluation_id}")

    return OptionResponse(
        id=option_id,
        evaluation_id=evaluation_id,
        name=option.name,
        type=option.type,
        vendor_name=option.vendor_name,
        description=option.description,
    )


# ---------------------------------------------------------------------------
# 6. remove_option
# ---------------------------------------------------------------------------


async def remove_option(
    evaluation_id: str, option_id: str, sp: SharePointClient
) -> None:
    """Remove an option and all associated scores and TCO data.

    Args:
        evaluation_id: UUID of the evaluation.
        option_id: UUID of the option to remove.
        sp: SharePoint client instance.

    Raises:
        ValueError: If the option is not found.
    """
    eval_file = await _resolve_eval_file(evaluation_id, sp)
    file_bytes = await sp._download_file(eval_file)
    wb = load_workbook(io.BytesIO(file_bytes))

    # Remove from Options sheet
    if SHEET_OPTIONS in wb.sheetnames:
        ws = wb[SHEET_OPTIONS]
        rows_to_delete = []
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=False), start=2):
            if row[0].value == option_id:
                rows_to_delete.append(row_idx)
        if not rows_to_delete:
            wb.close()
            raise ValueError(f"Option {option_id} not found in evaluation {evaluation_id}")
        for row_idx in reversed(rows_to_delete):
            ws.delete_rows(row_idx)

    # Remove associated scores
    if SHEET_SCORES in wb.sheetnames:
        ws = wb[SHEET_SCORES]
        rows_to_delete = []
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=False), start=2):
            if row[0].value == option_id:
                rows_to_delete.append(row_idx)
        for row_idx in reversed(rows_to_delete):
            ws.delete_rows(row_idx)

    # Remove associated TCO data
    if SHEET_TCO in wb.sheetnames:
        ws = wb[SHEET_TCO]
        rows_to_delete = []
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=False), start=2):
            if row[0].value == option_id:
                rows_to_delete.append(row_idx)
        for row_idx in reversed(rows_to_delete):
            ws.delete_rows(row_idx)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    await sp._upload_file(eval_file, buffer.getvalue())
    wb.close()

    logger.info(f"Removed option {option_id} from evaluation {evaluation_id}")



# ---------------------------------------------------------------------------
# 7. submit_scores
# ---------------------------------------------------------------------------


async def submit_scores(
    evaluation_id: str,
    option_id: str,
    scores: dict[str, int],
    source: str,
    sp: SharePointClient,
) -> ScoreResponse:
    """Submit dimension scores for an option.

    Validates each score is 1-5, stores in the Scores sheet, calculates
    weighted total, and reports unscored dimensions.

    Args:
        evaluation_id: UUID of the evaluation.
        option_id: UUID of the option being scored.
        scores: Dict mapping dimension keys to integer scores (1-5).
        source: Score source identifier (e.g. "manual", "slm_suggested").
        sp: SharePoint client instance.

    Returns:
        ScoreResponse with scoring results.

    Raises:
        ValueError: If any score is invalid.
    """
    # Validate scores
    for dimension, score in scores.items():
        if not scoring_matrix.validate_score(score):
            raise ValueError(
                f"Invalid score {score} for dimension '{dimension}': must be 1-5"
            )

    eval_file = await _resolve_eval_file(evaluation_id, sp)
    file_bytes = await sp._download_file(eval_file)
    wb = load_workbook(io.BytesIO(file_bytes))

    # Get or create Scores sheet
    if SHEET_SCORES in wb.sheetnames:
        ws = wb[SHEET_SCORES]
    else:
        ws = wb.create_sheet(SHEET_SCORES)
        ws.append(["OptionId", "Dimension", "Score", "Source"])

    # Remove existing scores for this option (to allow re-submission)
    rows_to_delete = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=False), start=2):
        if row[0].value == option_id:
            rows_to_delete.append(row_idx)
    for row_idx in reversed(rows_to_delete):
        ws.delete_rows(row_idx)

    # Append new scores
    for dimension, score in scores.items():
        ws.append([option_id, dimension, score, source])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    await sp._upload_file(eval_file, buffer.getvalue())
    wb.close()

    # Calculate weighted total
    weighted_total = scoring_matrix.calculate_weighted_total(scores)
    unscored = scoring_matrix.get_unscored_dimensions(scores)
    is_complete = len(unscored) == 0

    # Build response
    dimension_scores = [
        DimensionScoreResponse(
            dimension=dim,
            score=s,
            source=source,
            confidence=None,
        )
        for dim, s in scores.items()
    ]

    logger.info(
        f"Submitted scores for option {option_id} in evaluation {evaluation_id}: "
        f"weighted_total={weighted_total:.2f}, complete={is_complete}"
    )

    return ScoreResponse(
        option_id=option_id,
        scores=dimension_scores,
        weighted_total=weighted_total,
        is_complete=is_complete,
        unscored_dimensions=unscored,
    )


# ---------------------------------------------------------------------------
# 8. submit_tco
# ---------------------------------------------------------------------------


async def submit_tco(
    evaluation_id: str,
    option_id: str,
    tco_input: TCOSubmitRequest,
    sp: SharePointClient,
) -> TCOResultResponse:
    """Submit TCO data for an option and calculate metrics.

    Converts Pydantic input to tco_calculator.TCOInput, computes results,
    and stores the data in the TCO sheet.

    Args:
        evaluation_id: UUID of the evaluation.
        option_id: UUID of the option.
        tco_input: TCO submission request with cost data.
        sp: SharePoint client instance.

    Returns:
        TCOResultResponse with computed TCO metrics.
    """
    # Convert Pydantic models to tco_calculator dataclasses
    year_0 = tco_calculator.Year0Costs(
        license=tco_input.year_0.license,
        implementation=tco_input.year_0.implementation,
        migration=tco_input.year_0.migration,
        infrastructure=tco_input.year_0.infrastructure,
        training=tco_input.year_0.training,
        professional_services=tco_input.year_0.professional_services,
    )

    yearly_costs: list[tco_calculator.RecurringCosts] = []
    if tco_input.yearly_costs:
        for yc in tco_input.yearly_costs:
            yearly_costs.append(
                tco_calculator.RecurringCosts(
                    license_renewal=yc.license_renewal,
                    cloud=yc.cloud,
                    infrastructure=yc.infrastructure,
                    storage=yc.storage,
                    ktlo_people=yc.ktlo_people,
                    vendor_amc=yc.vendor_amc,
                    change_requests=yc.change_requests,
                    upgrades=yc.upgrades,
                    support=yc.support,
                )
            )

    calc_input = tco_calculator.TCOInput(
        option_id=option_id,
        year_0=year_0,
        yearly_costs=yearly_costs,
        user_count=tco_input.user_count,
    )

    # Calculate TCO
    tco_result = tco_calculator.calculate_tco(calc_input)

    # Store in TCO sheet
    eval_file = await _resolve_eval_file(evaluation_id, sp)
    file_bytes = await sp._download_file(eval_file)
    wb = load_workbook(io.BytesIO(file_bytes))

    if SHEET_TCO in wb.sheetnames:
        ws = wb[SHEET_TCO]
    else:
        ws = wb.create_sheet(SHEET_TCO)
        ws.append(["OptionId", "Category", "Year", "Amount"])

    # Remove existing TCO for this option
    rows_to_delete = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=False), start=2):
        if row[0].value == option_id:
            rows_to_delete.append(row_idx)
    for row_idx in reversed(rows_to_delete):
        ws.delete_rows(row_idx)

    # Store Year 0 costs
    ws.append([option_id, "year_0_total", 0, tco_result.year_0_total])
    ws.append([option_id, "three_year_tco", 0, tco_result.three_year_tco])
    ws.append([option_id, "five_year_tco", 0, tco_result.five_year_tco])
    ws.append([option_id, "annual_run_cost", 0, tco_result.annual_run_cost])
    if tco_result.cost_per_user is not None:
        ws.append([option_id, "cost_per_user", 0, tco_result.cost_per_user])

    # Store yearly totals
    for year_idx, yearly_total in enumerate(tco_result.yearly_totals):
        ws.append([option_id, "yearly_total", year_idx, yearly_total])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    await sp._upload_file(eval_file, buffer.getvalue())
    wb.close()

    logger.info(
        f"Submitted TCO for option {option_id} in evaluation {evaluation_id}: "
        f"5Y TCO={tco_result.five_year_tco:.2f}"
    )

    return TCOResultResponse(
        option_id=option_id,
        year_0_total=tco_result.year_0_total,
        three_year_tco=tco_result.three_year_tco,
        five_year_tco=tco_result.five_year_tco,
        annual_run_cost=tco_result.annual_run_cost,
        cost_per_user=tco_result.cost_per_user,
    )


# ---------------------------------------------------------------------------
# 9. determine_status
# ---------------------------------------------------------------------------


async def determine_status(
    evaluation_id: str, sp: SharePointClient
) -> EvaluationStatus:
    """Determine the lifecycle status of an evaluation based on data completeness.

    Logic:
    - No options → Draft
    - Options but incomplete scores/TCO → In_Progress
    - All options have complete scores and TCO → Scoring_Complete

    Args:
        evaluation_id: UUID of the evaluation.
        sp: SharePoint client instance.

    Returns:
        Computed EvaluationStatus.
    """
    data = await load_evaluation(evaluation_id, sp)
    options = data.get("options", [])
    scores = data.get("scores", [])
    tco_entries = data.get("tco", [])

    if not options:
        return EvaluationStatus.DRAFT

    # Check completeness for each option
    option_ids = [opt.get("OptionId", opt.get("optionid", "")) for opt in options]

    for opt_id in option_ids:
        # Check if option has complete scores (all 8 dimensions)
        opt_scores = [s for s in scores if s.get("OptionId") == opt_id]
        scored_dimensions = {s.get("Dimension") for s in opt_scores}
        if len(scored_dimensions) < len(scoring_matrix.DIMENSIONS):
            return EvaluationStatus.IN_PROGRESS

        # Check if option has TCO data
        opt_tco = [t for t in tco_entries if t.get("OptionId") == opt_id]
        if not opt_tco:
            return EvaluationStatus.IN_PROGRESS

    return EvaluationStatus.SCORING_COMPLETE
